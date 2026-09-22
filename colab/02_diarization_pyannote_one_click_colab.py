# ============================================================
# COLAB NOTEBOOK 2 — ONE-CLICK SPEAKER DIARIZATION WITH PYANNOTE
# ============================================================
#
# Purpose:
#   - Mount Google Drive.
#   - Find *_segmentos.json files produced by Notebook 1.
#   - Find the matching original audio file.
#   - Convert the audio to clean WAV 16 kHz mono.
#   - Run pyannote speaker diarization.
#   - Align diarization intervals with Whisper transcript segments.
#   - Export speaker-labelled transcripts and samples.
#
# Important design choice:
#   The notebook installs pyannote but does NOT import pyannote in the
#   active Colab kernel. Instead, it writes a runner script and launches it
#   in a fresh Python subprocess. This avoids runtime dependency conflicts
#   seen in Colab after package installation.
#
# Required Colab Secret:
#   TOKEN = Hugging Face token with access to pyannote model
#
# Designed for Google Colab.
# Paste this file into a Colab cell or convert it to a notebook.
# ============================================================

import os
import sys
import subprocess
from pathlib import Path

from google.colab import drive, userdata

# ------------------------------------------------------------
# CONFIGURATION
# ------------------------------------------------------------

DRIVE_FOLDER = Path("/content/drive/MyDrive/Transcrições e atas")
HF_TOKEN_SECRET_NAME = "TOKEN"

# Empty string means automatic speaker count.
# Set to "2", "3", "4", etc. if the number of speakers is known.
NUM_SPEAKERS = ""

# ------------------------------------------------------------
# NOTEBOOK-LEVEL COMMAND RUNNER
# ------------------------------------------------------------

def run(cmd, title):
    print("\n" + "=" * 80)
    print(title)
    print(" ".join(str(x) for x in cmd))
    print("=" * 80)

    process = subprocess.Popen(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        bufsize=1,
    )

    for line in process.stdout:
        print(line, end="", flush=True)

    code = process.wait()

    if code != 0:
        raise RuntimeError(f"Command failed: {title} | exit code {code}")

# ------------------------------------------------------------
# DRIVE + TOKEN
# ------------------------------------------------------------

print("\n" + "=" * 80)
print("ONE-CLICK SPEAKER DIARIZATION")
print("=" * 80)

if not Path("/content/drive/MyDrive").exists():
    drive.mount("/content/drive")
else:
    print("Drive already mounted.")

if not DRIVE_FOLDER.exists():
    raise RuntimeError(f"Drive folder not found: {DRIVE_FOLDER}")

HF_TOKEN = userdata.get(HF_TOKEN_SECRET_NAME)

if not HF_TOKEN:
    raise RuntimeError(f"Colab Secret not found: {HF_TOKEN_SECRET_NAME}")

HF_TOKEN = HF_TOKEN.strip()

if not HF_TOKEN.startswith("hf_"):
    raise RuntimeError("Invalid Hugging Face token. It should start with 'hf_'.")

os.environ["TOKEN_HF"] = HF_TOKEN
os.environ["NUM_SPEAKERS"] = NUM_SPEAKERS
os.environ["DRIVE_FOLDER"] = str(DRIVE_FOLDER)

print("Hugging Face token found in Colab Secrets.")
print("Drive folder:", DRIVE_FOLDER)

# ------------------------------------------------------------
# INSTALL MINIMAL DEPENDENCIES
# ------------------------------------------------------------

run(["apt-get", "update", "-qq"], "Update apt")
run(["apt-get", "install", "-y", "ffmpeg"], "Install ffmpeg")

run([
    sys.executable, "-m", "pip", "install",
    "-q", "--upgrade",
    "pip", "setuptools", "wheel",
], "Upgrade pip tooling")

run([
    sys.executable, "-m", "pip", "install",
    "-q", "--upgrade",
    "pyannote.audio",
    "python-docx",
    "huggingface_hub",
], "Install diarization dependencies")

# ------------------------------------------------------------
# WRITE EXTERNAL RUNNER SCRIPT
# ------------------------------------------------------------

runner_path = Path("/content/diarization_runner_one_click.py")

runner_path.write_text(r'''
import os
import json
import time
import shutil
import subprocess
from pathlib import Path
from collections import defaultdict
from datetime import datetime

import torch
from pyannote.audio import Pipeline
from huggingface_hub import whoami, hf_hub_download
from docx import Document

DRIVE_FOLDER = Path(os.environ["DRIVE_FOLDER"])
HF_TOKEN = os.environ["TOKEN_HF"]
NUM_SPEAKERS_ENV = os.environ.get("NUM_SPEAKERS", "").strip()
NUM_SPEAKERS = int(NUM_SPEAKERS_ENV) if NUM_SPEAKERS_ENV else None

DIARIZATION_MODEL = "pyannote/speaker-diarization-community-1"

AUDIO_EXTENSIONS = [
    ".m4a", ".mp3", ".wav", ".aac", ".flac",
    ".mp4", ".mov", ".mkv", ".webm", ".ogg",
]


def title(text):
    print("\n" + "=" * 80)
    print(text)
    print("=" * 80, flush=True)


def safe_name(name):
    invalid = '<>:"/\\|?*'
    for char in invalid:
        name = name.replace(char, "_")
    return name.strip()


def time_txt(seconds):
    seconds = max(0, int(float(seconds or 0)))
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def time_srt(seconds):
    total_ms = max(0, int(round(float(seconds or 0) * 1000)))
    h = total_ms // 3_600_000
    total_ms %= 3_600_000
    m = total_ms // 60_000
    total_ms %= 60_000
    s = total_ms // 1000
    ms = total_ms % 1000
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def find_segment_jsons():
    items = []

    for json_path in DRIVE_FOLDER.rglob("*_segmentos.json"):
        if json_path.name.endswith("_segmentos_oradores.json"):
            continue

        base = json_path.name.replace("_segmentos.json", "")
        output_marker = json_path.parent / f"{base}_segmentos_oradores.json"

        if output_marker.exists():
            continue

        items.append(json_path)

    return sorted(items, key=lambda p: p.stat().st_mtime)


def find_audio(base_name, transcript_data):
    drive_audio_path = transcript_data.get("caminho_audio_drive")

    if drive_audio_path:
        path = Path(drive_audio_path)
        if path.exists():
            return path

    audio_file_name = transcript_data.get("ficheiro_audio")

    if audio_file_name:
        direct = DRIVE_FOLDER / audio_file_name
        if direct.exists():
            return direct

        for candidate in DRIVE_FOLDER.rglob(audio_file_name):
            if candidate.exists():
                return candidate

    for ext in AUDIO_EXTENSIONS:
        direct = DRIVE_FOLDER / f"{base_name}{ext}"
        if direct.exists():
            return direct

    normalized_base = base_name.lower().replace("_", " ").strip()

    for candidate in DRIVE_FOLDER.rglob("*"):
        if candidate.is_file() and candidate.suffix.lower() in AUDIO_EXTENSIONS:
            stem = candidate.stem.lower().replace("_", " ").strip()
            if normalized_base in stem or stem in normalized_base:
                return candidate

    return None


def convert_to_wav(audio_path):
    base = safe_name(audio_path.stem)
    wav_path = Path("/content") / f"{base}_diarization_16k.wav"

    print("\nConverting audio to WAV 16 kHz mono...")
    print("Source:", audio_path)
    print("Output:", wav_path)

    cmd = [
        "ffmpeg",
        "-y",
        "-i", str(audio_path),
        "-vn",
        "-ac", "1",
        "-ar", "16000",
        "-sample_fmt", "s16",
        "-fflags", "+genpts",
        "-avoid_negative_ts", "make_zero",
        str(wav_path),
    ]

    process = subprocess.run(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )

    if process.returncode != 0:
        print(process.stdout)
        raise RuntimeError("ffmpeg audio conversion failed.")

    return wav_path


def overlap(a_start, a_end, b_start, b_end):
    return max(0.0, min(a_end, b_end) - max(a_start, b_start))


def assign_speaker(segment, diarization):
    scores = defaultdict(float)

    for interval in diarization:
        value = overlap(
            segment["inicio"],
            segment["fim"],
            interval["inicio"],
            interval["fim"],
        )

        if value > 0:
            scores[interval["orador"]] += value

    if not scores:
        return "ORADOR_DESCONHECIDO"

    best_speaker, best_score = max(scores.items(), key=lambda item: item[1])

    if best_score < 0.05:
        return "ORADOR_DESCONHECIDO"

    return best_speaker


def write_txt(segments, path, audio_name):
    speakers = sorted(set(s["orador"] for s in segments))

    with open(path, "w", encoding="utf-8") as f:
        f.write(f"SPEAKER-LABELLED TRANSCRIPT — {audio_name}\n")
        f.write(f"Detected speakers: {len(speakers)}\n")
        f.write("=" * 80 + "\n\n")

        for s in segments:
            f.write(
                f"[{time_txt(s['inicio'])} - {time_txt(s['fim'])}] "
                f"{s['orador']}: {s['texto']}\n\n"
            )


def write_srt(segments, path):
    with open(path, "w", encoding="utf-8") as f:
        n = 1

        for s in segments:
            text = s["texto"].strip()
            if not text:
                continue

            f.write(f"{n}\n")
            f.write(f"{time_srt(s['inicio'])} --> {time_srt(s['fim'])}\n")
            f.write(f"{s['orador']}: {text}\n\n")
            n += 1


def write_docx(segments, path, audio_name):
    doc = Document()
    doc.add_heading("Speaker-labelled transcript", level=1)
    doc.add_paragraph(f"Source file: {audio_name}")
    doc.add_paragraph("")

    last_speaker = None

    for s in segments:
        speaker = s["orador"]
        text = s["texto"].strip()

        if not text:
            continue

        if speaker != last_speaker:
            p = doc.add_paragraph()
            r = p.add_run(speaker)
            r.bold = True
            last_speaker = speaker

        p = doc.add_paragraph()
        p.add_run(f"[{time_txt(s['inicio'])} - {time_txt(s['fim'])}] ").italic = True
        p.add_run(text)

    doc.save(path)


def write_speaker_samples(segments, path, max_samples=8):
    samples = defaultdict(list)

    # Prefer longer fragments spread through the file.
    by_speaker = defaultdict(list)
    fallback = defaultdict(list)

    for s in segments:
        speaker = s["orador"]
        text = s["texto"].strip()
        if not text:
            continue

        fallback[speaker].append(s)

        if len(text) >= 45:
            by_speaker[speaker].append(s)

    for speaker in sorted(fallback.keys()):
        candidates = by_speaker.get(speaker) or fallback.get(speaker) or []
        if not candidates:
            continue

        n = min(max_samples, len(candidates))

        if n == 1:
            indexes = [len(candidates) // 2]
        else:
            indexes = sorted(set(round(i * (len(candidates) - 1) / (n - 1)) for i in range(n)))

        samples[speaker] = [candidates[i] for i in indexes]

    with open(path, "w", encoding="utf-8") as f:
        f.write("SAMPLES FOR SPEAKER IDENTIFICATION\n")
        f.write("=" * 80 + "\n\n")

        for speaker in sorted(samples.keys()):
            f.write(f"{speaker}\n")
            f.write("-" * 80 + "\n")

            for s in samples[speaker]:
                f.write(f"- [{time_txt(s['inicio'])}] {s['texto']}\n")

            f.write("\n")


def save_json(data, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# ------------------------------------------------------------
# MAIN
# ------------------------------------------------------------

title("DIARIZATION RUNNER")

print("Drive folder:", DRIVE_FOLDER)
print("Torch:", torch.__version__)
print("CUDA:", torch.cuda.is_available())

device = torch.device("cuda" if torch.cuda.is_available() else "cpu")

if torch.cuda.is_available():
    print("GPU:", torch.cuda.get_device_name(0))
else:
    print("No GPU detected. Diarization will be slower.")

print("Hugging Face account:", whoami(token=HF_TOKEN).get("name"))

hf_hub_download(
    repo_id=DIARIZATION_MODEL,
    filename="config.yaml",
    token=HF_TOKEN,
)

print("\nLoading pyannote pipeline...")

pipeline = Pipeline.from_pretrained(
    DIARIZATION_MODEL,
    token=HF_TOKEN,
)

pipeline.to(device)

print("Pipeline loaded on:", device)

json_files = find_segment_jsons()

if not json_files:
    print("No new *_segmentos.json files found for diarization.")
    raise SystemExit(0)

print("Transcript JSON files found:", len(json_files))
for path in json_files:
    print("-", path)

for json_path in json_files:
    title(f"PROCESSING: {json_path.name}")
    start_time = time.time()

    with open(json_path, "r", encoding="utf-8") as f:
        transcript_data = json.load(f)

    base_name = json_path.name.replace("_segmentos.json", "")
    audio_path = find_audio(base_name, transcript_data)

    if not audio_path:
        print("Audio not found for:", base_name)
        print("Skipping file.")
        continue

    raw_segments = transcript_data.get("segmentos", [])
    segments = []

    for i, s in enumerate(raw_segments, start=1):
        text = str(s.get("texto", "")).strip()
        if not text:
            continue

        start = float(s.get("inicio", 0))
        end = float(s.get("fim", start))

        if end <= start:
            end = start + 0.1

        segments.append({
            "id": int(s.get("id", i)),
            "inicio": start,
            "fim": end,
            "texto": text,
        })

    if not segments:
        print("No valid transcript segments found.")
        continue

    print("Audio:", audio_path)
    print("Transcript segments:", len(segments))
    print("Configured speaker count:", NUM_SPEAKERS if NUM_SPEAKERS else "automatic")

    wav_path = convert_to_wav(audio_path)

    print("\nRunning diarization...")

    if NUM_SPEAKERS is None:
        result = pipeline(str(wav_path))
    else:
        result = pipeline(str(wav_path), num_speakers=NUM_SPEAKERS)

    print("Diarization finished. Converting result...")

    if hasattr(result, "exclusive_speaker_diarization"):
        annotation = result.exclusive_speaker_diarization
    elif hasattr(result, "speaker_diarization"):
        annotation = result.speaker_diarization
    else:
        annotation = result

    diarization = []

    for turn, _, speaker in annotation.itertracks(yield_label=True):
        diarization.append({
            "inicio": float(turn.start),
            "fim": float(turn.end),
            "orador": str(speaker),
        })

    print("Speech intervals detected:", len(diarization))

    if not diarization:
        print("No speaker intervals returned by diarization model.")
        continue

    speaker_segments = []

    for s in segments:
        speaker_segments.append({
            "id": s["id"],
            "inicio": s["inicio"],
            "fim": s["fim"],
            "orador": assign_speaker(s, diarization),
            "texto": s["texto"],
        })

    speakers = sorted(set(s["orador"] for s in speaker_segments))
    print("Detected speakers:", ", ".join(speakers))

    output_folder = json_path.parent
    base = json_path.name.replace("_segmentos.json", "")

    json_out = output_folder / f"{base}_segmentos_oradores.json"
    txt_out = output_folder / f"{base}_transcricao_oradores.txt"
    docx_out = output_folder / f"{base}_transcricao_oradores.docx"
    srt_out = output_folder / f"{base}_legendas_oradores.srt"
    samples_out = output_folder / f"{base}_amostras_oradores.txt"
    zip_out = output_folder / f"{base}_COM_ORADORES.zip"

    output_data = {
        "ficheiro_audio": audio_path.name,
        "ficheiro_wav_diarizacao": str(wav_path),
        "data_processamento_diarizacao": datetime.now().isoformat(),
        "modelo_diarizacao": DIARIZATION_MODEL,
        "num_speakers_configurado": NUM_SPEAKERS,
        "oradores_detetados": speakers,
        "diarizacao_bruta": diarization,
        "segmentos": speaker_segments,
        "transcricao_original": transcript_data,
    }

    print("Exporting JSON...")
    save_json(output_data, json_out)

    print("Exporting TXT...")
    write_txt(speaker_segments, txt_out, audio_path.name)

    print("Exporting DOCX...")
    write_docx(speaker_segments, docx_out, audio_path.name)

    print("Exporting SRT...")
    write_srt(speaker_segments, srt_out)

    print("Exporting speaker samples...")
    write_speaker_samples(speaker_segments, samples_out)

    print("Creating ZIP package...")
    if zip_out.exists():
        zip_out.unlink()

    temp_zip_base = Path("/content") / f"{base}_COM_ORADORES"
    temp_zip = shutil.make_archive(str(temp_zip_base), "zip", root_dir=output_folder)
    shutil.copy2(temp_zip, zip_out)

    print("\nCREATED:")
    print("JSON:    ", json_out)
    print("TXT:     ", txt_out)
    print("DOCX:    ", docx_out)
    print("SRT:     ", srt_out)
    print("Samples: ", samples_out)
    print("ZIP:     ", zip_out)
    print("Elapsed:", round((time.time() - start_time) / 60, 1), "minutes")

title("DIARIZATION FINISHED")
print("Check the same Drive result folders created by Notebook 1.")
''', encoding="utf-8")

print("Runner created:", runner_path)

# ------------------------------------------------------------
# RUN EXTERNAL SCRIPT IN A FRESH PYTHON PROCESS
# ------------------------------------------------------------

print("\nRunning diarization in a fresh Python subprocess...\n")

env = os.environ.copy()

process = subprocess.Popen(
    [sys.executable, str(runner_path)],
    stdout=subprocess.PIPE,
    stderr=subprocess.STDOUT,
    text=True,
    bufsize=1,
    env=env,
)

for line in process.stdout:
    print(line, end="", flush=True)

code = process.wait()

if code != 0:
    raise RuntimeError(f"Diarization failed with exit code {code}")

print("\nONE-CLICK DIARIZATION FINISHED.")
