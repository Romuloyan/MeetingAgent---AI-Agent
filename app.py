import json
import os
import time
import hashlib
from pathlib import Path
from collections import defaultdict
from datetime import datetime
from io import BytesIO

import streamlit as st
from docx import Document
from pypdf import PdfReader
from dotenv import load_dotenv
from google import genai


PASTA_DEFAULT = r"C:\Users\Pc\OneDrive\Agente IA\Transcricoes_atas"
MODELO_GEMINI_DEFAULT = "gemini-3.5-flash-lite"


def tempo_txt(segundos):
    segundos = max(0, int(float(segundos or 0)))
    h = segundos // 3600
    m = (segundos % 3600) // 60
    s = segundos % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def carregar_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def guardar_json(path, dados):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)


def guardar_txt(path, texto):
    with open(path, "w", encoding="utf-8") as f:
        f.write(texto)


def procurar_reunioes(pasta_base):
    pasta = Path(pasta_base)
    if not pasta.exists():
        return []
    return sorted(pasta.rglob("*_segmentos_oradores.json"))


def obter_segmentos(dados):
    segmentos = dados.get("segmentos", [])
    limpos = []

    for i, s in enumerate(segmentos, start=1):
        texto = str(s.get("texto", "")).strip()
        if not texto:
            continue

        inicio = float(s.get("inicio", 0))
        fim = float(s.get("fim", inicio))
        if fim <= inicio:
            fim = inicio + 0.1

        limpos.append({
            "id": int(s.get("id", i)),
            "inicio": inicio,
            "fim": fim,
            "orador": str(s.get("orador", "SEM_ORADOR")),
            "texto": texto,
        })

    return limpos


def obter_oradores(segmentos):
    oradores = sorted(set(s["orador"] for s in segmentos))
    if "ORADOR_DESCONHECIDO" in oradores:
        oradores.remove("ORADOR_DESCONHECIDO")
        oradores.append("ORADOR_DESCONHECIDO")
    return oradores


def escolher_amostras(segmentos, max_por_orador=8):
    por_orador = defaultdict(list)
    fallback = defaultdict(list)

    for s in segmentos:
        orador = s["orador"]
        texto = s["texto"].strip()
        if texto:
            fallback[orador].append(s)
        if len(texto) >= 45:
            por_orador[orador].append(s)

    resultado = {}

    for orador in sorted(fallback.keys()):
        candidatos = por_orador.get(orador) or fallback.get(orador) or []
        if not candidatos:
            resultado[orador] = []
            continue

        n = min(max_por_orador, len(candidatos))
        if n == 1:
            indices = [len(candidatos) // 2]
        else:
            indices = sorted(set(round(i * (len(candidatos) - 1) / (n - 1)) for i in range(n)))

        resultado[orador] = [candidatos[i] for i in indices]

    return resultado


def escrever_txt_com_nomes(segmentos, path, titulo):
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"TRANSCRIÇÃO COM NOMES — {titulo}\n")
        f.write("=" * 80 + "\n\n")

        ultimo = None
        for s in segmentos:
            orador = s["orador"]
            texto = s["texto"].strip()
            if orador != ultimo:
                f.write("\n")
                f.write(f"{orador}\n")
                f.write("-" * 80 + "\n")
                ultimo = orador
            f.write(f"[{tempo_txt(s['inicio'])} - {tempo_txt(s['fim'])}] {texto}\n\n")


def escrever_docx_com_nomes(segmentos, path, titulo):
    doc = Document()
    doc.add_heading("Transcrição com nomes", level=1)
    doc.add_paragraph(f"Reunião: {titulo}")
    doc.add_paragraph("")

    ultimo = None
    for s in segmentos:
        orador = s["orador"]
        texto = s["texto"].strip()
        if not texto:
            continue
        if orador != ultimo:
            p = doc.add_paragraph()
            r = p.add_run(orador)
            r.bold = True
            ultimo = orador
        p = doc.add_paragraph()
        p.add_run(f"[{tempo_txt(s['inicio'])} - {tempo_txt(s['fim'])}] ").italic = True
        p.add_run(texto)

    doc.save(path)


def aplicar_mapa_oradores(json_path, dados, segmentos, mapa):
    base = json_path.name.replace("_segmentos_oradores.json", "")
    pasta = json_path.parent
    segmentos_com_nomes = []

    for s in segmentos:
        orador_original = s["orador"]
        nome_real = mapa.get(orador_original, "").strip() or orador_original
        novo = dict(s)
        novo["orador_original"] = orador_original
        novo["orador"] = nome_real
        segmentos_com_nomes.append(novo)

    dados_out = dict(dados)
    dados_out["mapa_oradores"] = mapa
    dados_out["data_atribuicao_oradores"] = datetime.now().isoformat()
    dados_out["segmentos"] = segmentos_com_nomes

    json_out = pasta / f"{base}_segmentos_com_nomes.json"
    txt_out = pasta / f"{base}_transcricao_com_nomes.txt"
    docx_out = pasta / f"{base}_transcricao_com_nomes.docx"
    mapa_out = pasta / f"{base}_mapa_oradores.json"

    guardar_json(json_out, dados_out)
    guardar_json(mapa_out, mapa)
    escrever_txt_com_nomes(segmentos_com_nomes, txt_out, base)
    escrever_docx_com_nomes(segmentos_com_nomes, docx_out, base)

    return {"json": json_out, "txt": txt_out, "docx": docx_out, "mapa": mapa_out}


def chave_curta(texto):
    return hashlib.md5(texto.encode("utf-8")).hexdigest()[:10]


def carregar_mapa_existente(json_path):
    base = json_path.name.replace("_segmentos_oradores.json", "")
    mapa_path = json_path.parent / f"{base}_mapa_oradores.json"
    if mapa_path.exists():
        try:
            return carregar_json(mapa_path)
        except Exception:
            return {}
    return {}


def ler_docx_bytes(data):
    doc = Document(BytesIO(data))
    partes = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            partes.append(txt)
    return "\n".join(partes)


def ler_pdf_bytes(data):
    reader = PdfReader(BytesIO(data))
    partes = []
    for i, page in enumerate(reader.pages, start=1):
        texto = (page.extract_text() or "").strip()
        if texto:
            partes.append(f"\n--- PÁGINA {i} ---\n{texto}")
    return "\n".join(partes)


def ler_txt_bytes(data):
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="ignore")


def ler_upload(uploaded_file):
    nome = uploaded_file.name
    data = uploaded_file.getvalue()
    ext = Path(nome).suffix.lower()

    if ext == ".txt":
        texto = ler_txt_bytes(data)
    elif ext == ".docx":
        texto = ler_docx_bytes(data)
    elif ext == ".pdf":
        texto = ler_pdf_bytes(data)
    else:
        texto = ""

    return {"nome": nome, "texto": texto.strip()}


def segmentos_para_texto(segmentos):
    return "\n".join(
        f"[{tempo_txt(s['inicio'])} - {tempo_txt(s['fim'])}] {s['orador']}: {s['texto']}"
        for s in segmentos
    )


def dividir_texto(texto, max_chars=90000):
    linhas = texto.splitlines()
    blocos = []
    atual = []
    tamanho = 0

    for linha in linhas:
        if tamanho + len(linha) > max_chars and atual:
            blocos.append("\n".join(atual))
            atual = []
            tamanho = 0
        atual.append(linha)
        tamanho += len(linha) + 1

    if atual:
        blocos.append("\n".join(atual))

    return blocos


def obter_cliente_gemini():
    load_dotenv()
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY não encontrada no ficheiro .env")
    return genai.Client(api_key=api_key)


def chamar_gemini(modelo, prompt, tentativas=3):
    cliente = obter_cliente_gemini()

    modelos = []
    modelo = (modelo or "").strip()
    if modelo:
        modelos.append(modelo)

    for fallback in ["gemini-3.5-flash-lite", "gemini-3.6-flash"]:
        if fallback not in modelos:
            modelos.append(fallback)

    esperas = [10, 25, 45]
    ultimo_erro = None

    for modelo_atual in modelos:
        st.info(f"A usar modelo: {modelo_atual}")
        for i in range(tentativas):
            try:
                resposta = cliente.models.generate_content(model=modelo_atual, contents=prompt)
                texto = resposta.text or ""
                if texto.strip():
                    return texto.strip()
                ultimo_erro = RuntimeError("Resposta vazia do Gemini.")
            except Exception as e:
                ultimo_erro = e
                erro_txt = str(e)
                if "503" in erro_txt or "UNAVAILABLE" in erro_txt or "high demand" in erro_txt:
                    espera = esperas[min(i, len(esperas) - 1)]
                    st.warning(f"{modelo_atual} falhou por sobrecarga na tentativa {i + 1}/{tentativas}. A esperar {espera}s.")
                    time.sleep(espera)
                    continue
                st.warning(f"{modelo_atual} falhou na tentativa {i + 1}/{tentativas}: {erro_txt[:250]}")
                time.sleep(5)
        st.warning(f"A trocar de modelo. {modelo_atual} não respondeu de forma estável.")

    raise RuntimeError(f"Gemini falhou em todos os modelos. Último erro: {ultimo_erro}")


def escrever_docx_markdown_simples(path, titulo, texto):
    doc = Document()
    doc.add_heading(titulo, level=1)

    for linha in texto.splitlines():
        l = linha.strip()
        if not l:
            doc.add_paragraph("")
        elif l.startswith("# "):
            doc.add_heading(l.replace("# ", "").strip(), level=1)
        elif l.startswith("## "):
            doc.add_heading(l.replace("## ", "").strip(), level=2)
        elif l.startswith("### "):
            doc.add_heading(l.replace("### ", "").strip(), level=3)
        elif l.startswith("- "):
            doc.add_paragraph(l[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(l)

    doc.save(path)


def criar_prompt_extracao_bloco(nome_reuniao, bloco_num, bloco_total, texto_bloco):
    return f"""
És um assistente de atas em português europeu.

Vais analisar APENAS um bloco de transcrição de uma reunião.
Não inventes. Não acrescentes factos externos.
Se algo não estiver claro, escreve "não confirmado".

Reunião: {nome_reuniao}
Bloco: {bloco_num}/{bloco_total}

Extrai:
1. Temas tratados neste bloco
2. Decisões ou deliberações mencionadas
3. Tarefas / responsáveis / prazos, se existirem
4. Pendentes
5. Pontos de conflito ou dúvidas
6. Frases relevantes para ata

Transcrição do bloco:
{texto_bloco}
"""


def criar_prompt_ata(nome_reuniao, texto_base, contexto_documentos, instrucoes_extra):
    return f"""
És um secretário de reunião e vais produzir uma ata preliminar em português europeu.

REGRAS OBRIGATÓRIAS:
- Usa apenas a transcrição e os documentos fornecidos.
- Não inventes decisões, presenças, votações, datas ou responsáveis.
- Se algo não estiver explícito, escreve "não confirmado" ou "não ficou claro".
- A transcrição é a fonte principal do que foi dito.
- Documentos opcionais, ordem de trabalhos, notas ou ata preliminar servem apenas como contexto.
- Usa linguagem formal, clara e adequada a ata.
- Mantém neutralidade.

INSTRUÇÕES DO UTILIZADOR:
{instrucoes_extra}

DOCUMENTOS / CONTEXTO OPCIONAL:
{contexto_documentos}

MATERIAL BASE DA REUNIÃO:
{texto_base}

PRODUZ A RESPOSTA COM ESTA ESTRUTURA:

# ATA PRELIMINAR

## 1. Identificação da reunião
- Reunião:
- Data:
- Local / modalidade:
- Participantes:
- Observações sobre dados não confirmados:

## 2. Ordem de trabalhos / temas tratados

## 3. Síntese da reunião

## 4. Pontos discutidos
Organiza por tópico.

## 5. Decisões / deliberações
Só incluir decisões se estiverem suportadas na transcrição.

## 6. Tarefas e responsáveis
Formato:
- Tarefa:
- Responsável:
- Prazo:
- Estado / observações:

## 7. Pendentes

## 8. Riscos, dúvidas ou pontos a confirmar

## 9. Encerramento

# NOTAS DE FIABILIDADE
Indica aqui limitações da ata, dúvidas de transcrição, speakers incertos ou informação que precisa de confirmação humana.
"""


def gerar_ata_gemini(json_com_nomes_path, modelo, documentos, notas_livres, instrucoes_extra):
    dados = carregar_json(json_com_nomes_path)
    segmentos = obter_segmentos(dados)
    nome_reuniao = json_com_nomes_path.name.replace("_segmentos_com_nomes.json", "")
    texto_transcricao = segmentos_para_texto(segmentos)

    partes_contexto = []
    for doc in documentos:
        if doc["texto"]:
            partes_contexto.append(f"\n\n===== DOCUMENTO: {doc['nome']} =====\n{doc['texto']}")

    if notas_livres.strip():
        partes_contexto.append(f"\n\n===== NOTAS LIVRES DO UTILIZADOR =====\n{notas_livres.strip()}")

    contexto_documentos = "\n".join(partes_contexto).strip()
    blocos = dividir_texto(texto_transcricao, max_chars=90000)

    if len(blocos) == 1:
        texto_base = texto_transcricao
        analise_intermedia = "Transcrição enviada diretamente para geração da ata."
    else:
        resumos = []
        progresso = st.progress(0)
        for i, bloco in enumerate(blocos, start=1):
            st.write(f"A analisar bloco {i}/{len(blocos)}...")
            prompt_bloco = criar_prompt_extracao_bloco(nome_reuniao, i, len(blocos), bloco)
            resumo = chamar_gemini(modelo, prompt_bloco)
            resumos.append(f"\n\n===== BLOCO {i}/{len(blocos)} =====\n{resumo}")
            progresso.progress(i / len(blocos))
        analise_intermedia = "\n".join(resumos)
        texto_base = analise_intermedia

    prompt_ata = criar_prompt_ata(nome_reuniao, texto_base, contexto_documentos, instrucoes_extra)
    st.write("A gerar ata final...")
    ata = chamar_gemini(modelo, prompt_ata)

    prompt_tarefas = f"""
A partir da ata abaixo, extrai apenas uma lista objetiva de tarefas, responsáveis, prazos e pendentes.
Não inventes.
Se não houver responsável ou prazo, escreve "não confirmado".

ATA:
{ata}
"""

    st.write("A extrair tarefas e pendentes...")
    tarefas = chamar_gemini(modelo, prompt_tarefas)

    pasta = json_com_nomes_path.parent
    base = json_com_nomes_path.name.replace("_segmentos_com_nomes.json", "")

    ata_txt = pasta / f"{base}_ata_preliminar_gemini.txt"
    ata_docx = pasta / f"{base}_ata_preliminar_gemini.docx"
    analise_txt = pasta / f"{base}_analise_intermedia_gemini.txt"
    tarefas_txt = pasta / f"{base}_tarefas_pendentes_gemini.txt"

    guardar_txt(ata_txt, ata)
    guardar_txt(analise_txt, analise_intermedia)
    guardar_txt(tarefas_txt, tarefas)
    escrever_docx_markdown_simples(ata_docx, "Ata preliminar", ata)

    return {"ata_txt": ata_txt, "ata_docx": ata_docx, "analise_txt": analise_txt, "tarefas_txt": tarefas_txt}


st.set_page_config(page_title="MeetingAgent", layout="wide")
st.title("MeetingAgent — Identificação de Oradores e Ata")
st.write("Fase 1: identificar SPEAKERs. Fase 2: carregar contexto opcional e gerar ata preliminar com Gemini.")

pasta_base = st.text_input("Pasta onde estão os resultados do Colab", value=PASTA_DEFAULT)

if not Path(pasta_base).exists():
    st.error("A pasta não existe. Confirma o caminho.")
    st.stop()

reunioes = procurar_reunioes(pasta_base)
if not reunioes:
    st.warning("Não encontrei ficheiros *_segmentos_oradores.json nessa pasta.")
    st.stop()

st.success(f"Reuniões diarizadas encontradas: {len(reunioes)}")
opcoes = [str(p) for p in reunioes]

selecionada = st.selectbox("Escolhe a reunião", options=opcoes, format_func=lambda x: Path(x).name)
json_path = Path(selecionada)
dados = carregar_json(json_path)
segmentos = obter_segmentos(dados)
oradores = obter_oradores(segmentos)
amostras = escolher_amostras(segmentos)
mapa_existente = carregar_mapa_existente(json_path)

base = json_path.name.replace("_segmentos_oradores.json", "")
json_com_nomes_path = json_path.parent / f"{base}_segmentos_com_nomes.json"

st.divider()
col1, col2, col3 = st.columns(3)
with col1:
    st.metric("Segmentos", len(segmentos))
with col2:
    st.metric("Oradores detetados", len(oradores))
with col3:
    duracao = max([s["fim"] for s in segmentos], default=0)
    st.metric("Duração aprox.", tempo_txt(duracao))

st.header("Fase 1 — Atribuição de nomes")
st.write("Lê as frases de cada SPEAKER e escreve o nome real correspondente.")

mapa = {}
chave_reuniao = chave_curta(str(json_path))

for orador in oradores:
    with st.expander(orador, expanded=False):
        lista = amostras.get(orador, [])
        if not lista:
            st.write("Sem amostras úteis.")
        else:
            for s in lista:
                st.markdown(f"**[{tempo_txt(s['inicio'])}]** {s['texto']}")

        valor_default = mapa_existente.get(orador, "")
        if not valor_default and orador == "ORADOR_DESCONHECIDO":
            valor_default = "ORADOR_DESCONHECIDO"

        nome = st.text_input(f"{orador} =", value=valor_default, key=f"{chave_reuniao}_{orador}")
        mapa[orador] = nome.strip()

st.subheader("Gerar ficheiros com nomes reais")
if st.button("Gerar transcrição com nomes", type="primary"):
    saidas = aplicar_mapa_oradores(json_path, dados, segmentos, mapa)
    st.success("Ficheiros criados com sucesso.")
    for _, path in saidas.items():
        st.code(str(path))

st.divider()
st.header("Fase 2 — Gerar ata com Gemini")

if not json_com_nomes_path.exists():
    st.warning("Ainda não existe *_segmentos_com_nomes.json. Primeiro gera a transcrição com nomes na Fase 1.")
    st.stop()

st.success("Transcrição com nomes encontrada. Já podes gerar ata.")

modelo = st.text_input("Modelo Gemini", value=MODELO_GEMINI_DEFAULT)

uploads = st.file_uploader(
    "Documentos opcionais da reunião: ordem de trabalhos, notas, ata preliminar, anexos",
    type=["txt", "docx", "pdf"],
    accept_multiple_files=True,
)

notas_livres = st.text_area(
    "Notas livres opcionais",
    height=120,
    placeholder="Exemplo: reunião da direção; usar tom formal; confirmar nomes; não inventar decisões...",
)

instrucoes_extra = st.text_area(
    "Instruções para a ata",
    value=(
        "Gerar uma ata preliminar formal em português europeu. "
        "Não inventar decisões. Separar decisões, tarefas, pendentes e pontos a confirmar. "
        "Quando houver dúvidas de transcrição, assinalar como não confirmado."
    ),
    height=100,
)

if st.button("Gerar ata preliminar com Gemini", type="primary"):
    try:
        documentos = []
        if uploads:
            for u in uploads:
                with st.spinner(f"A ler {u.name}..."):
                    documentos.append(ler_upload(u))

        saidas_ata = gerar_ata_gemini(
            json_com_nomes_path=json_com_nomes_path,
            modelo=modelo,
            documentos=documentos,
            notas_livres=notas_livres,
            instrucoes_extra=instrucoes_extra,
        )

        st.success("Ata preliminar criada com sucesso.")
        st.write("Ficheiros criados:")
        for _, path in saidas_ata.items():
            st.code(str(path))

    except Exception as e:
        st.error("Falhou ao gerar ata.")
        st.exception(e)
