# ============================================================
# COLAB NOTEBOOK 1 — TRANSCRIPTION WITH FASTER-WHISPER
# ============================================================
#
# Purpose:
#   - Mount Google Drive.
#   - Look for new audio/video files in a Drive folder.
#   - Transcribe them with faster-whisper.
#   - Export timestamped TXT, plain TXT, SRT, JSON segments, DOCX, and ZIP.
#
# Designed for Google Colab.
# Paste this file into a Colab cell or convert it to a notebook.
#
# Real recordings and generated outputs are intentionally not included
# in the GitHub repository.
# ============================================================

import json
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

# ------------------------------------------------------------
# INSTALL DEPENDENCIES
# ------------------------------------------------------------

subprocess.check_call(["apt-get", "update", "-qq"])
subprocess.check_call(["apt-get", "install", "-y", "ffmpeg"])
subprocess.check_call([
    sys.executable, "-m", "pip", "install",
    "-q", "--upgrade",
    "faster-whisper",
    "python-docx",
])

# ------------------------------------------------------------
# IMPORTS AFTER INSTALL
# ------------------------------------------------------------

import torch
from faster_whisper import WhisperModel
from google.colab import drive
from docx import Document

# ------------------------------------------------------------
# CONFIGURATION
# ------------------------------------------------------------

DRIVE_FOLDER = Path("/content/drive/MyDrive/Transcrições e atas")
OUTPUT_ROOT_NAME = "_resultados_transcricao"

MODEL_NAME = "large-v3"
LANGUAGE = "pt"

SUPPORTED_EXTENSIONS = {
    ".m4a", ".mp3", ".wav", ".aac", ".flac",
    ".mp4", ".mov", ".mkv", ".webm", ".ogg",
}

# ------------------------------------------------------------
# HELPERS
# ------------------------------------------------------------

def print_title(text):
    print("\n" + "=" * 80)
    print(text)
    print("=" * 80, flush=True)


def safe_name(name):
    invalid = '<>:"/\\|?*'
    for char in invalid:
        name = name.replace(char, "_")
    return name.strip()


def time_txt(seconds):
    seconds = max(0, int(float(seconds or 0)))
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def time_srt(seconds):
    total_ms = max(0, int(round(float(seconds or 0) * 1000)))
    h = total_ms // 3_600_000
    total_ms %= 3_600_000
    m = total_ms // 60_000
    total_ms %= 60_000
    s = total_ms // 1000
    ms = total_ms % 1000
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def find_input_files(folder):
    files = []
    for path in folder.iterdir():
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            files.append(path)
    return sorted(files, key=lambda p: p.stat().st_mtime)


def load_processed_marker(output_folder, base_name):
    return output_folder / safe_name(base_name) / f"{safe_name(base_name)}_segmentos.json"


def export_timestamp_txt(segments, path):
    with open(path, "w", encoding="utf-8") as f:
        for s in segments:
            f.write(f"[{time_txt(s['inicio'])} - {time_txt(s['fim'])}] {s['texto']}\n\n")


def export_plain_txt(segments, path):
    text = " ".join(s["texto"].strip() for s in segments if s["texto"].strip())
    with open(path, "w", encoding="utf-8") as f:
        f.write(text.strip() + "\n")


def export_srt(segments, path):
    with open(path, "w", encoding="utf-8") as f:
        n = 1
        for s in segments:
            text = s["texto"].strip()
            if not text:
                continue
            f.write(f"{n}\n")
            f.write(f"{time_srt(s['inicio'])} --> {time_srt(s['fim'])}\n")
            f.write(f"{text}\n\n")
            n += 1


def export_docx(segments, path, audio_name):
    doc = Document()
    doc.add_heading("Transcription", level=1)
    doc.add_paragraph(f"Source file: {audio_name}")
    doc.add_paragraph("")

    for s in segments:
        p = doc.add_paragraph()
        p.add_run(f"[{time_txt(s['inicio'])} - {time_txt(s['fim'])}] ").italic = True
        p.add_run(s["texto"].strip())

    doc.save(path)


def save_json(data, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

# ------------------------------------------------------------
# DRIVE
# ------------------------------------------------------------

print_title("MOUNTING GOOGLE DRIVE")

if not Path("/content/drive/MyDrive").exists():
    drive.mount("/content/drive")
else:
    print("Drive already mounted.")

if not DRIVE_FOLDER.exists():
    raise RuntimeError(f"Drive folder not found: {DRIVE_FOLDER}")

OUTPUT_ROOT = DRIVE_FOLDER / OUTPUT_ROOT_NAME
OUTPUT_ROOT.mkdir(exist_ok=True)

print("Input folder:", DRIVE_FOLDER)
print("Output folder:", OUTPUT_ROOT)

# ------------------------------------------------------------
# DEVICE / MODEL
# ------------------------------------------------------------

print_title("LOADING WHISPER MODEL")

if torch.cuda.is_available():
    DEVICE = "cuda"
    COMPUTE_TYPE = "float16"
    print("CUDA available:", torch.cuda.get_device_name(0))
else:
    DEVICE = "cpu"
    COMPUTE_TYPE = "int8"
    print("CUDA not available. Running on CPU.")

print("Model:", MODEL_NAME)
print("Language:", LANGUAGE)
print("Device:", DEVICE)
print("Compute type:", COMPUTE_TYPE)

model = WhisperModel(
    MODEL_NAME,
    device=DEVICE,
    compute_type=COMPUTE_TYPE,
)

# ------------------------------------------------------------
# PROCESS FILES
# ------------------------------------------------------------

input_files = find_input_files(DRIVE_FOLDER)

if not input_files:
    print("No audio/video files found.")
    raise SystemExit(0)

print_title("FILES FOUND")
for f in input_files:
    print("-", f.name)

for audio_path in input_files:
    base = safe_name(audio_path.stem)
    meeting_output = OUTPUT_ROOT / base
    meeting_output.mkdir(exist_ok=True)

    marker = meeting_output / f"{base}_segmentos.json"
    if marker.exists():
        print_title(f"SKIPPING ALREADY PROCESSED: {audio_path.name}")
        print("Existing marker:", marker)
        continue

    print_title(f"TRANSCRIBING: {audio_path.name}")

    segments_iter, info = model.transcribe(
        str(audio_path),
        language=LANGUAGE,
        beam_size=5,
        vad_filter=True,
        vad_parameters={"min_silence_duration_ms": 500},
        condition_on_previous_text=True,
    )

    segments = []

    for i, seg in enumerate(segments_iter, start=1):
        item = {
            "id": i,
            "inicio": float(seg.start),
            "fim": float(seg.end),
            "texto": seg.text.strip(),
        }
        segments.append(item)
        print(f"[{time_txt(item['inicio'])} - {time_txt(item['fim'])}] {item['texto']}", flush=True)

    if not segments:
        print("No transcription segments returned.")
        continue

    json_data = {
        "ficheiro_audio": audio_path.name,
        "caminho_audio_drive": str(audio_path),
        "data_processamento": datetime.now().isoformat(),
        "modelo": MODEL_NAME,
        "idioma": LANGUAGE,
        "device": DEVICE,
        "compute_type": COMPUTE_TYPE,
        "duracao_estimativa_segundos": max(s["fim"] for s in segments),
        "segmentos": segments,
    }

    timestamp_txt = meeting_output / f"{base}_transcricao_timestamp.txt"
    plain_txt = meeting_output / f"{base}_transcricao_corrida.txt"
    srt_path = meeting_output / f"{base}_legendas.srt"
    json_path = meeting_output / f"{base}_segmentos.json"
    docx_path = meeting_output / f"{base}_transcricao.docx"
    zip_path = meeting_output / f"{base}_TRANSCRICAO_COMPLETA.zip"

    print("\nExporting files...")
    export_timestamp_txt(segments, timestamp_txt)
    export_plain_txt(segments, plain_txt)
    export_srt(segments, srt_path)
    save_json(json_data, json_path)
    export_docx(segments, docx_path, audio_path.name)

    if zip_path.exists():
        zip_path.unlink()

    temp_zip_base = Path("/content") / f"{base}_TRANSCRICAO_COMPLETA"
    temp_zip = shutil.make_archive(str(temp_zip_base), "zip", root_dir=meeting_output)
    shutil.copy2(temp_zip, zip_path)

    print("\nCREATED:")
    print("Timestamp TXT:", timestamp_txt)
    print("Plain TXT:    ", plain_txt)
    print("SRT:          ", srt_path)
    print("JSON:         ", json_path)
    print("DOCX:         ", docx_path)
    print("ZIP:          ", zip_path)

print_title("TRANSCRIPTION FINISHED")
